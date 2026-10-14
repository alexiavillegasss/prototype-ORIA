import openpyxl

wb_new = openpyxl.load_workbook('config/rules/tableauORIA - NOMENCLATURE-CASSI.xlsx', data_only=True)
ws_new = wb_new['02_Besoins_x_structures (2)']

print("Unique CASSI prestations per detaille:")
mapping = {}
for r in range(2, ws_new.max_row+1):
    prob = ws_new.cell(r, 2).value
    besoin = ws_new.cell(r, 3).value
    prest = ws_new.cell(r, 4).value
    detail = ws_new.cell(r, 5).value
    kw = ws_new.cell(r, 6).value
    if detail not in mapping:
        mapping[detail] = []
    mapping[detail].append((prob, besoin, prest, kw))

for detail, items in mapping.items():
    print(f"\n[{detail}] ({len(items)} variations):")
    for prob, bes, prest, kw in items:
        if prest and prest != 'x':
            print(f"   -> Prestation: {prest} | CASSI: {prob} > {bes} | KW: {kw}")
        else:
            print(f"   -> (Sans prestation CASSI spécifique) | KW: {kw}")
